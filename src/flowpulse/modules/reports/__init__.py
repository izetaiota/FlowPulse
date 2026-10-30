"""报表中心（P5-T5.4）：把站内数据导出为 Excel（openpyxl，懒加载可选依赖）。

支持六类报表：xhs-monitors / xhs-contents / douyin-monitors / douyin-contents /
downloads / accounts。文件名 flowpulse-<kind>-<日期>.xlsx。
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response

router = APIRouter(prefix="/api/v1/reports", tags=["reports"])

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def get_reports_service():
    raise RuntimeError("Dependency override is not configured")


class ReportsService:
    def __init__(self, xhs_monitor_repo, douyin_monitor_repo, download_store, account_repo):
        self.xhs = xhs_monitor_repo
        self.douyin = douyin_monitor_repo
        self.downloads = download_store
        self.accounts = account_repo

    # ---- 数据装配（返回 (表头, 行)） ----

    def _monitor_rows(self, repo) -> tuple[list[str], list[list]]:
        stats = repo.content_stats()
        monitors = repo.list_monitors()
        names = {m["id"]: (m.get("name") or m["target"]) for m in monitors}
        headers = ["ID", "类型", "名称", "目标", "周期(分钟)", "作品数", "状态", "最近扫描", "扫描状态", "最近错误"]
        rows = [
            [
                m["id"],
                "创作者" if m["kind"] == "creator" else "关键词",
                m["name"],
                m["target"],
                m["interval_minutes"],
                stats.get(m["id"], 0),
                "监控中" if m["status"] == "active" else "已暂停",
                m.get("last_scan_at") or "",
                m.get("last_scan_status") or "",
                (m.get("last_error") or "")[:200],
            ]
            for m in monitors
        ]
        return headers, rows

    def _content_rows(self, repo, platform: str) -> tuple[list[str], list[list]]:
        monitors = repo.list_monitors()
        names = {m["id"]: (m.get("name") or m["target"]) for m in monitors}
        headers = ["平台", "发现时间", "来源监控", "标题", "作者", "作者ID", "类型", "点赞", "链接"]
        rows = []
        contents, _ = repo.list_contents(page=1, page_size=5000)
        for c in contents:
            rows.append(
                [
                    platform,
                    c.get("first_seen_at") or "",
                    names.get(c.get("monitor_id"), ""),
                    c.get("title") or "",
                    c.get("author") or "",
                    c.get("author_id") or c.get("author_id_xhs", "") or "",
                    "视频" if (c.get("media_type") or c.get("note_type")) == "video" else "图文",
                    str(c.get("liked_count") or ""),
                    c.get("url") or "",
                ]
            )
        return headers, rows

    def _download_rows(self) -> tuple[list[str], list[list]]:
        headers = ["创建时间", "来源", "链接", "标题", "状态", "产物文件", "大小(MB)", "错误"]
        rows = []
        for t in self.downloads.list():
            size_mb = round((t.file_size or 0) / 1024 / 1024, 2)
            rows.append(
                [
                    t.created_at.isoformat(),
                    t.source,
                    t.url,
                    t.title or "",
                    t.status,
                    t.filename or "",
                    size_mb,
                    (t.error or "")[:200],
                ]
            )
        return headers, rows

    def _account_rows(self) -> tuple[list[str], list[list]]:
        headers = ["平台", "账号标识", "名称", "类型", "状态", "最近体检", "体检结果", "备注"]
        rows = []
        for a in self.accounts.list_accounts():
            rows.append(
                [
                    a["platform"],
                    a["account_key"],
                    a["name"],
                    "模块账号" if a["kind"] == "builtin" else "手动登记",
                    a["status"],
                    a.get("last_check_at") or "",
                    "通过" if a.get("last_check_ok") else ("失败" if a.get("last_check_ok") is False else ""),
                    (a.get("last_check_message") or "")[:200],
                ]
            )
        return headers, rows

    # ---- xlsx 生成 ----

    def build(self, kind: str) -> tuple[bytes, str]:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RuntimeError("未安装 openpyxl（pip install openpyxl），无法导出报表") from exc

        builders = {
            "xhs-monitors": lambda: self._monitor_rows(self.xhs),
            "xhs-contents": lambda: self._content_rows(self.xhs, "小红书"),
            "douyin-monitors": lambda: self._monitor_rows(self.douyin),
            "douyin-contents": lambda: self._content_rows(self.douyin, "抖音"),
            "downloads": self._download_rows,
            "accounts": self._account_rows,
        }
        if kind not in builders:
            raise KeyError(f"未知报表：{kind}")
        headers, rows = builders[kind]()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = kind[:31]
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        # 表头加粗 + 简单列宽（按表头长度估）
        from openpyxl.styles import Font

        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for index, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(
                12, min(50, len(str(header)) * 2.4 + 6)
            )
        buffer = io.BytesIO()
        workbook.save(buffer)
        date_str = datetime.now(timezone.utc).strftime("%Y%m%d")
        return buffer.getvalue(), f"flowpulse-{kind}-{date_str}.xlsx"


REPORT_CATALOG: list[dict] = [
    {"kind": "xhs-monitors", "label": "小红书监控目标", "desc": "全部监控目标与扫描状态"},
    {"kind": "xhs-contents", "label": "小红书作品记录", "desc": "监控发现的全部作品明细"},
    {"kind": "douyin-monitors", "label": "抖音监控目标", "desc": "全部监控目标与扫描状态"},
    {"kind": "douyin-contents", "label": "抖音作品记录", "desc": "监控发现的全部作品明细"},
    {"kind": "downloads", "label": "下载任务历史", "desc": "链接下载任务与产物清单"},
    {"kind": "accounts", "label": "账号清单", "desc": "统一账号中心全部账号与体检状态"},
]


@router.get("")
def list_reports() -> dict:
    return {"reports": REPORT_CATALOG}


@router.get("/{kind}.xlsx")
def export_report(kind: str, service=Depends(get_reports_service)) -> Response:
    try:
        content, filename = service.build(kind)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return Response(
        content=content,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
